from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook


WIDTH, HEIGHT = 832, 594
BACKGROUND = "#1A1E43"
BLUE = "#0070C0"
YELLOW = "#FFC000"


def read_data(workbook_path: Path) -> tuple[list[str], list[float]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["2 带均值柱形图"]
    categories = [str(worksheet.cell(row=row, column=2).value) for row in range(3, 9)]
    values = [float(worksheet.cell(row=row, column=3).value) for row in range(3, 9)]
    workbook.close()
    return categories, values


def reproduce(workbook_path: Path, output_path: Path) -> None:
    categories, values = read_data(workbook_path)
    average = float(np.mean(values))
    regular = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")

    figure = plt.figure(
        figsize=(WIDTH / 100, HEIGHT / 100), dpi=100, facecolor=BACKGROUND
    )
    axes = figure.add_axes(
        [49 / WIDTH, (HEIGHT - 495) / HEIGHT, (802 - 49) / WIDTH, 355 / HEIGHT],
        facecolor=BACKGROUND,
    )
    x = np.arange(len(categories))
    axes.bar(x, values, width=0.315, color=BLUE, edgecolor="none", zorder=2)
    axes.set_xlim(-0.5, 5.5)
    axes.set_ylim(0, 4000)
    axes.set_xticks(x, categories)
    axes.set_yticks([])
    axes.tick_params(axis="x", colors="white", length=0, pad=10, labelsize=12)
    for label in axes.get_xticklabels():
        label.set_fontproperties(regular)
    for side in ("left", "right", "top"):
        axes.spines[side].set_visible(False)
    axes.spines["bottom"].set_color("#86889A")
    axes.spines["bottom"].set_linewidth(1.5)

    for index, value in enumerate(values):
        inside = index == len(values) - 1
        axes.text(
            index,
            value - 135 if inside else value + 125,
            f"{value:.0f}",
            ha="center",
            va="center" if inside else "bottom",
            color="white",
            fontsize=10.5,
            fontproperties=regular,
            zorder=4,
        )

    # 均值线从首根柱左边缘延伸至末根柱中心，与 Excel 原图一致。
    axes.hlines(average, -0.15, 5.02, colors=YELLOW, linewidth=2.0, zorder=3)
    axes.text(
        4.48,
        average + 175,
        f"平均值：{average:.0f}",
        color=YELLOW,
        fontsize=11,
        fontproperties=regular,
        ha="left",
        va="bottom",
    )

    figure.text(
        56 / WIDTH,
        1 - 34 / HEIGHT,
        "3月各区域销量分布",
        color="white",
        fontsize=28,
        fontproperties=bold,
        ha="left",
        va="top",
    )
    figure.text(
        55 / WIDTH,
        1 - 91 / HEIGHT,
        "东北销量最多占比总销量的22%，华南销量最低",
        color="white",
        fontsize=20,
        fontproperties=regular,
        ha="left",
        va="top",
    )
    figure.text(
        52 / WIDTH,
        1 - 558 / HEIGHT,
        "*注：数据来源于公司销售系统，统计日期截至2022.03.31",
        color="white",
        fontsize=11,
        fontproperties=regular,
        ha="left",
        va="top",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=100, facecolor=BACKGROUND, edgecolor="none")
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第02图：带均值柱形图")
    parser.add_argument(
        "--workbook", type=Path, default=directory / "第二章 图表(前15).xlsx"
    )
    parser.add_argument(
        "--output", type=Path, default=directory / "02_带均值柱形图_Python复现.png"
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    reproduce(args.workbook.resolve(), args.output.resolve())
    print(f"已生成：{args.output.resolve()}")
