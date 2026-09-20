from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Rectangle
from openpyxl import load_workbook


WIDTH, HEIGHT = 832, 588
BACKGROUND = "#1A1E43"
BLUE = "#007CC3"
PINK = "#E94C6A"


def read_data(workbook_path: Path) -> tuple[list[str], list[float], list[float]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["6 蝴蝶图"]
    rows = [
        (
            str(worksheet.cell(row=row, column=2).value),
            float(worksheet.cell(row=row, column=4).value),
            float(worksheet.cell(row=row, column=6).value),
        )
        for row in range(3, 8)
    ]
    workbook.close()
    rows.reverse()
    return [r[0] for r in rows], [r[1] for r in rows], [r[2] for r in rows]


def reproduce(workbook_path: Path, output_path: Path) -> None:
    categories, current, previous = read_data(workbook_path)
    regular = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")
    figure = plt.figure(figsize=(WIDTH / 100, HEIGHT / 100), dpi=100, facecolor=BACKGROUND)
    scale = 220 / max(current)
    left_anchor = 368
    right_anchor = 488
    bar_height = 28
    row_tops = [189, 259, 329, 400, 470]

    for category, left_value, right_value, top in zip(categories, current, previous, row_tops):
        left_width = left_value * scale
        right_width = right_value * scale
        figure.add_artist(Rectangle(((left_anchor - left_width) / WIDTH, (HEIGHT - top - bar_height) / HEIGHT), left_width / WIDTH, bar_height / HEIGHT, transform=figure.transFigure, facecolor=BLUE, edgecolor="none"))
        figure.add_artist(Rectangle((right_anchor / WIDTH, (HEIGHT - top - bar_height) / HEIGHT), right_width / WIDTH, bar_height / HEIGHT, transform=figure.transFigure, facecolor=PINK, edgecolor="none"))
        center_y = (HEIGHT - top - bar_height / 2) / HEIGHT
        figure.text((left_anchor - left_width + 12) / WIDTH, center_y, f"{left_value:.0f}", color="white", fontsize=10.5, fontproperties=regular, ha="left", va="center")
        figure.text((right_anchor + right_width - 12) / WIDTH, center_y, f"{right_value:.0f}", color="white", fontsize=10.5, fontproperties=regular, ha="right", va="center")
        figure.text(428 / WIDTH, center_y, category, color="white", fontsize=12, fontproperties=regular, ha="center", va="center")

    figure.text(54 / WIDTH, 1 - 47 / HEIGHT, "2022年上半年各区域对比去年销量", color="white", fontsize=27, fontproperties=bold, ha="left", va="top")
    figure.text(54 / WIDTH, 1 - 100 / HEIGHT, "2022年整体销量高于2021年，只有东北区域较2021有所下降", color="white", fontsize=18, fontproperties=regular, ha="left", va="top")
    figure.text(308 / WIDTH, 1 - 156 / HEIGHT, "2022", color="#0085CE", fontsize=14, fontproperties=regular, ha="center", va="top")
    figure.text(518 / WIDTH, 1 - 156 / HEIGHT, "2021", color=PINK, fontsize=14, fontproperties=regular, ha="center", va="top")
    figure.text(52 / WIDTH, 1 - 550 / HEIGHT, "*注：数据来源于公司销售系统，统计日期截至2022.06.30", color="white", fontsize=11, fontproperties=regular, ha="left", va="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=100, facecolor=BACKGROUND, edgecolor="none")
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第06图：蝴蝶图")
    parser.add_argument("--workbook", type=Path, default=directory / "第二章 图表(前15).xlsx")
    parser.add_argument("--output", type=Path, default=directory / "06_蝴蝶图_Python复现.png")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    reproduce(args.workbook.resolve(), args.output.resolve())
    print(f"已生成：{args.output.resolve()}")
