from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook


CANVAS_WIDTH = 832
CANVAS_HEIGHT = 617
BACKGROUND = "#1A1E43"
GRID_COLOR = "#404361"
BAR_TOP = "#0070C0"
BAR_BOTTOM = "#00B0F0"


def read_chart_data(workbook_path: Path) -> tuple[list[str], list[float]]:
    """读取“1 渐变柱形图”工作表中的分类和销售量。"""
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["1 渐变柱形图"]
    categories = [worksheet.cell(row=row, column=2).value for row in range(3, 9)]
    values = [worksheet.cell(row=row, column=3).value for row in range(3, 9)]
    workbook.close()
    return [str(value) for value in categories], [float(value) for value in values]


def add_gradient_bar(
    axes: plt.Axes,
    center: float,
    height: float,
    width: float,
    gradient: LinearSegmentedColormap,
) -> None:
    """绘制与 Excel 一致的蓝—青纵向渐变柱。"""
    vertical_gradient = np.linspace(0.0, 1.0, 512).reshape(512, 1)
    axes.imshow(
        vertical_gradient,
        extent=(center - width / 2, center + width / 2, 0, height),
        origin="upper",
        aspect="auto",
        cmap=gradient,
        interpolation="bicubic",
        zorder=3,
    )


def reproduce_chart(workbook_path: Path, output_path: Path) -> None:
    categories, values = read_chart_data(workbook_path)
    regular_font = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold_font = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")

    figure = plt.figure(
        figsize=(CANVAS_WIDTH / 100, CANVAS_HEIGHT / 100),
        dpi=100,
        facecolor=BACKGROUND,
    )

    # Excel 原图绘图区：左 120 px、右 752 px、上 197 px、下 521 px。
    axes = figure.add_axes(
        [
            120 / CANVAS_WIDTH,
            (CANVAS_HEIGHT - 521) / CANVAS_HEIGHT,
            (752 - 120) / CANVAS_WIDTH,
            (521 - 197) / CANVAS_HEIGHT,
        ],
        facecolor=BACKGROUND,
    )

    x_positions = np.arange(len(categories), dtype=float)
    gradient = LinearSegmentedColormap.from_list(
        "excel_blue_gradient", [BAR_TOP, BAR_BOTTOM]
    )
    for x_position, value in zip(x_positions, values):
        add_gradient_bar(axes, x_position, value, 0.315, gradient)
        axes.text(
            x_position,
            value + 135,
            f"{value:.0f}",
            ha="center",
            va="bottom",
            color="white",
            fontsize=10.5,
            fontproperties=regular_font,
            zorder=4,
        )

    axes.set_xlim(-0.5, len(categories) - 0.5)
    axes.set_ylim(0, 4000)
    axes.set_xticks(x_positions, categories)
    axes.set_yticks([0, 1000, 2000, 3000, 4000])
    axes.tick_params(axis="x", colors="white", labelsize=12, length=0, pad=10)
    axes.tick_params(axis="y", colors="white", labelsize=10.5, length=0, pad=12)
    for label in axes.get_xticklabels() + axes.get_yticklabels():
        label.set_fontproperties(regular_font)

    axes.grid(axis="y", color=GRID_COLOR, linewidth=1.35, linestyle=(0, (9, 4)))
    axes.set_axisbelow(True)
    for spine in axes.spines.values():
        spine.set_visible(False)

    figure.text(
        66 / CANVAS_WIDTH,
        1 - 45 / CANVAS_HEIGHT,
        "3月各区域销量分布",
        ha="left",
        va="top",
        color="white",
        fontsize=28,
        fontproperties=bold_font,
    )
    figure.text(
        66 / CANVAS_WIDTH,
        1 - 104 / CANVAS_HEIGHT,
        "东北销量最多占比总销量的22%，华南销量最低",
        ha="left",
        va="top",
        color="white",
        fontsize=20,
        fontproperties=regular_font,
    )
    figure.text(
        50 / CANVAS_WIDTH,
        1 - 580 / CANVAS_HEIGHT,
        "*注：数据来源于公司销售系统，统计日期截至2022.03.31",
        ha="left",
        va="top",
        color="white",
        fontsize=11,
        fontproperties=regular_font,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(
        output_path,
        dpi=100,
        facecolor=BACKGROUND,
        edgecolor="none",
        format="png",
    )
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第01图：渐变柱形图")
    parser.add_argument(
        "--workbook",
        type=Path,
        default=script_dir / "第二章 图表(前15).xlsx",
        help="源 Excel 文件路径",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "01_渐变柱形图_Python复现.png",
        help="输出 PNG 文件路径",
    )
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    reproduce_chart(arguments.workbook.resolve(), arguments.output.resolve())
    print(f"已生成：{arguments.output.resolve()}")
