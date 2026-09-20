from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from matplotlib.patches import FancyBboxPatch
from matplotlib.transforms import Affine2D
from openpyxl import load_workbook


WIDTH, HEIGHT = 832, 617
BACKGROUND = "#1A1E43"
GRID = "#404361"
BAR_BOTTOM = "#0058B8"
BAR_TOP = "#08E2E7"
LABEL_BACKGROUND = "#123B72"
LABEL_DOT = "#5AACE3"


def read_data(workbook_path: Path) -> tuple[list[str], list[float]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["3 渐变圆角柱形图"]
    categories = [str(worksheet.cell(row=row, column=2).value) for row in range(3, 9)]
    values = [float(worksheet.cell(row=row, column=3).value) for row in range(3, 9)]
    workbook.close()
    return categories, values


def rounded_gradient_bar(axes: plt.Axes, x: float, value: float) -> None:
    width = 0.17
    radius = width / 2
    clip = FancyBboxPatch(
        (x - width / 2, 0),
        width,
        value / (1200 / 329 * (657 / 6)),
        boxstyle=f"round,pad=0,rounding_size={radius}",
        linewidth=0,
        facecolor="none",
        transform=Affine2D().scale(1, 1200 / 329 * (657 / 6)) + axes.transData,
    )
    axes.add_patch(clip)
    gradient = np.linspace(0, 1, 512).reshape(512, 1)
    cmap = LinearSegmentedColormap.from_list("cyan_blue", [BAR_BOTTOM, BAR_TOP])
    image = axes.imshow(
        gradient,
        extent=(x - width / 2, x + width / 2, 0, value),
        origin="lower",
        aspect="auto",
        cmap=cmap,
        interpolation="bicubic",
        zorder=3,
    )
    image.set_clip_path(clip)


def reproduce(workbook_path: Path, output_path: Path) -> None:
    categories, values = read_data(workbook_path)
    regular = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")
    figure = plt.figure(
        figsize=(WIDTH / 100, HEIGHT / 100), dpi=100, facecolor=BACKGROUND
    )
    axes = figure.add_axes(
        [112 / WIDTH, (HEIGHT - 510) / HEIGHT, (769 - 112) / WIDTH, 329 / HEIGHT],
        facecolor=BACKGROUND,
    )
    x = np.arange(len(categories), dtype=float)
    for index, value in enumerate(values):
        rounded_gradient_bar(axes, index, value)
        marker_value = value + (120 if index == 0 else 85)
        axes.vlines(index, value - 5, marker_value, color="#0A5595", linewidth=2, zorder=4)
        axes.text(
            index + 0.12,
            marker_value,
            f"   {value:.0f}  ",
            ha="center",
            va="center",
            color="#D6E8F5",
            fontsize=11,
            fontproperties=regular,
            bbox=dict(boxstyle="round,pad=0.48", facecolor=LABEL_BACKGROUND, edgecolor="none"),
            zorder=5,
        )
        axes.scatter(
            [index - 0.01], [marker_value], s=62, color=LABEL_DOT, edgecolors="none", zorder=6
        )

    axes.set_xlim(-0.5, 5.5)
    axes.set_ylim(0, 1200)
    axes.set_xticks(x, categories)
    axes.set_yticks(np.arange(0, 1201, 200))
    axes.tick_params(axis="x", colors="white", labelsize=12, length=0, pad=10)
    axes.tick_params(axis="y", colors="white", labelsize=10.5, length=0, pad=12)
    for label in axes.get_xticklabels() + axes.get_yticklabels():
        label.set_fontproperties(regular)
        label.set_fontsize(12 if label in axes.get_xticklabels() else 10.5)
    axes.grid(axis="y", color=GRID, linewidth=1.0, linestyle=(0, (5, 3)))
    axes.set_axisbelow(True)
    for spine in axes.spines.values():
        spine.set_visible(False)

    figure.text(
        57 / WIDTH, 1 - 30 / HEIGHT, "3月商品销量对比",
        color="white", fontsize=28, fontproperties=bold, ha="left", va="top"
    )
    figure.text(
        57 / WIDTH, 1 - 89 / HEIGHT,
        "防晒销量最多，3月销量856；面膜最少，3月销量523",
        color="white", fontsize=20, fontproperties=regular, ha="left", va="top"
    )
    figure.text(
        50 / WIDTH, 1 - 579 / HEIGHT,
        "*注：数据来源于公司销售系统，统计日期截至2022.03.31",
        color="white", fontsize=11, fontproperties=regular, ha="left", va="top"
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=100, facecolor=BACKGROUND, edgecolor="none")
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第03图：渐变圆角柱形图")
    parser.add_argument("--workbook", type=Path, default=directory / "第二章 图表(前15).xlsx")
    parser.add_argument("--output", type=Path, default=directory / "03_渐变圆角柱形图_Python复现.png")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    reproduce(args.workbook.resolve(), args.output.resolve())
    print(f"已生成：{args.output.resolve()}")
