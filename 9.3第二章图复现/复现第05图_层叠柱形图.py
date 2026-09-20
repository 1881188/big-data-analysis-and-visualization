from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Rectangle
from openpyxl import load_workbook


WIDTH, HEIGHT = 832, 588
BACKGROUND = "#1A1E43"
BLUE = "#0070C0"
PINK = "#E74E69"


def read_data(workbook_path: Path) -> tuple[list[str], list[float], list[float]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["5 层叠柱形图"]
    categories = [str(worksheet.cell(row=row, column=2).value) for row in range(3, 9)]
    sales = [float(worksheet.cell(row=row, column=3).value) for row in range(3, 9)]
    profit = [float(worksheet.cell(row=row, column=4).value) for row in range(3, 9)]
    workbook.close()
    return categories, sales, profit


def reproduce(workbook_path: Path, output_path: Path) -> None:
    categories, sales, profit = read_data(workbook_path)
    regular = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")
    figure = plt.figure(figsize=(WIDTH / 100, HEIGHT / 100), dpi=100, facecolor=BACKGROUND)
    axes = figure.add_axes([73 / WIDTH, (HEIGHT - 497) / HEIGHT, 673 / WIDTH, 343 / HEIGHT], facecolor=BACKGROUND)
    x = np.arange(len(categories), dtype=float)
    bar_width = 0.26
    axes.bar(x - 0.09, sales, width=bar_width, color=BLUE, edgecolor="none", zorder=2)
    axes.bar(x + 0.09, profit, width=bar_width, color=PINK, edgecolor="none", zorder=3)
    axes.set_xlim(-0.5, 5.5)
    axes.set_ylim(0, 6000)
    axes.set_xticks(x, categories)
    axes.set_yticks([])
    axes.tick_params(axis="x", colors="white", labelsize=11, length=0, pad=11)
    for label in axes.get_xticklabels():
        label.set_fontproperties(regular)
    for side in ("left", "right", "top"):
        axes.spines[side].set_visible(False)
    axes.spines["bottom"].set_color("#454866")
    axes.spines["bottom"].set_linewidth(1.5)

    for index, value in enumerate(sales):
        axes.text(index - 0.09, value + 175, f"{value:.0f}", color="white", fontsize=10.5, fontproperties=regular, ha="center", va="bottom")
    for index, value in enumerate(profit):
        axes.text(index + 0.09, value + 170, f"{value:.0f}", color="white", fontsize=10.5, fontproperties=regular, ha="center", va="bottom")

    figure.text(48 / WIDTH, 1 - 51 / HEIGHT, "2021年至今季度销售额(万)和利润额(万)", color="white", fontsize=27, fontproperties=bold, ha="left", va="top")
    figure.text(48 / WIDTH, 1 - 114 / HEIGHT, "2022年第二季度销售额首次出现下降，降幅达到15%", color="white", fontsize=18, fontproperties=regular, ha="left", va="top")
    figure.text(40 / WIDTH, 1 - 550 / HEIGHT, "*注：数据来源于公司销售系统，统计日期截至2022.06.30", color="white", fontsize=11, fontproperties=regular, ha="left", va="top")

    legend_x, legend_y = 664 / WIDTH, (HEIGHT - 328) / HEIGHT
    figure.add_artist(Rectangle((legend_x, legend_y), 82 / WIDTH, 25 / HEIGHT, transform=figure.transFigure, facecolor=BACKGROUND, edgecolor=BLUE, linewidth=1.5))
    figure.text(705 / WIDTH, (HEIGHT - 315) / HEIGHT, "销售额", color="white", fontsize=10.5, fontproperties=regular, ha="center", va="center")
    legend_y2 = (HEIGHT - 372) / HEIGHT
    figure.add_artist(Rectangle((legend_x, legend_y2), 82 / WIDTH, 25 / HEIGHT, transform=figure.transFigure, facecolor=BACKGROUND, edgecolor=PINK, linewidth=1.5))
    figure.text(705 / WIDTH, (HEIGHT - 359) / HEIGHT, "利润额", color="white", fontsize=10.5, fontproperties=regular, ha="center", va="center")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=100, facecolor=BACKGROUND, edgecolor="none")
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第05图：层叠柱形图")
    parser.add_argument("--workbook", type=Path, default=directory / "第二章 图表(前15).xlsx")
    parser.add_argument("--output", type=Path, default=directory / "05_层叠柱形图_Python复现.png")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    reproduce(args.workbook.resolve(), args.output.resolve())
    print(f"已生成：{args.output.resolve()}")
