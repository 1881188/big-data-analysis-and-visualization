from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
from matplotlib.patches import Polygon


WIDTH, HEIGHT = 832, 616
BACKGROUND = "#1A1E43"
GRID = "#404361"
COLORS = ["#7BBDD5", "#49A098", "#E66B4C", "#FFC000", "#0097E0", "#0070C0", "#4A5BD1", "#464CAC"]


def read_data(workbook_path: Path) -> tuple[list[str], list[float]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    worksheet = workbook["4 标注柱形图"]
    categories = [str(worksheet.cell(row=row, column=2).value) for row in range(3, 11)]
    values = [float(worksheet.cell(row=row, column=3).value) for row in range(3, 11)]
    workbook.close()
    return categories, values


def reproduce(workbook_path: Path, output_path: Path) -> None:
    categories, values = read_data(workbook_path)
    regular = FontProperties(fname=r"C:\Windows\Fonts\msyh.ttc")
    bold = FontProperties(fname=r"C:\Windows\Fonts\msyhbd.ttc")
    figure = plt.figure(figsize=(WIDTH / 100, HEIGHT / 100), dpi=100, facecolor=BACKGROUND)
    axes = figure.add_axes([118 / WIDTH, (HEIGHT - 515) / HEIGHT, 656 / WIDTH, 330 / HEIGHT], facecolor=BACKGROUND)
    x = np.arange(len(categories), dtype=float)
    axes.bar(x, values, width=0.50, color=COLORS, edgecolor="none", zorder=3)
    axes.set_xlim(-0.5, 7.5)
    axes.set_ylim(0, 10000)
    axes.set_xticks(x, categories)
    axes.set_yticks(np.arange(0, 10001, 2000))
    axes.tick_params(axis="x", colors="white", labelsize=10.5, length=0, pad=10)
    axes.tick_params(axis="y", colors="white", labelsize=9.5, length=0, pad=10)
    for label in axes.get_xticklabels() + axes.get_yticklabels():
        label.set_fontproperties(regular)
        label.set_fontsize(12 if label in axes.get_xticklabels() else 10.5)
    axes.grid(axis="y", color=GRID, linewidth=1.0, linestyle=(0, (6, 3)))
    axes.set_axisbelow(True)
    for spine in axes.spines.values():
        spine.set_visible(False)

    for index, (value, color) in enumerate(zip(values, COLORS)):
        axes.text(
            index, value + 420, f"{value:.0f}", ha="center", va="bottom",
            color="white", fontsize=10.5, fontproperties=regular,
            bbox=dict(boxstyle="square,pad=0.40", facecolor=color, edgecolor="none"), zorder=4,
        )

        axes.add_patch(Polygon([(index-.16,value+420),(index-.04,value+420),(index,value+40)],closed=True,facecolor=color,edgecolor="none",zorder=4))

    figure.text(54 / WIDTH, 1 - 40 / HEIGHT, "2021年商品销量情况", color="white", fontsize=26, fontproperties=bold, ha="left", va="top")
    figure.text(61 / WIDTH, 1 - 100 / HEIGHT, "口红销量最好达9221，是眼影最低值2645近3.5倍", color="white", fontsize=17, fontproperties=regular, ha="left", va="top")
    figure.text(56 / WIDTH, 1 - 578 / HEIGHT, "*注：数据来源于公司销售系统，统计日期截至2022.08.31", color="white", fontsize=10.5, fontproperties=regular, ha="left", va="top")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=100, facecolor=BACKGROUND, edgecolor="none")
    plt.close(figure)


def parse_args() -> argparse.Namespace:
    directory = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="复现第二章第04图：标注柱形图")
    parser.add_argument("--workbook", type=Path, default=directory / "第二章 图表(前15).xlsx")
    parser.add_argument("--output", type=Path, default=directory / "04_标注柱形图_Python复现.png")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    reproduce(args.workbook.resolve(), args.output.resolve())
    print(f"已生成：{args.output.resolve()}")
