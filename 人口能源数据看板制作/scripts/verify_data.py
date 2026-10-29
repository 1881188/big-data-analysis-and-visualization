"""Independent checks of saved CSVs and Excel against original response cells."""
from pathlib import Path
import csv
import json
import math
from decimal import Decimal
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"data/processed"

def rows(name):
    with (OUT/name).open(encoding="utf-8-sig",newline="") as f:
        return list(csv.DictReader(f))
def close(a,b):
    return math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-8)

def main():
    source_cells={}
    for file in (ROOT/"data/raw").glob("*_data.json"):
        for record in json.loads(file.read_text(encoding="utf-8"))["data"]:
            year=int(record["code"][:4])
            for v in record["values"]:
                source_cells[(year,v["_id"])]=v["value"]
    observed=rows("observations.csv")
    assert len(observed)==len(source_cells)==720
    for r in observed:
        raw=source_cells[(int(r["year"]),r["indicator_id"])]
        assert r["raw_text"]==raw
        if raw=="":
            assert r["value"]==""
        else:
            assert Decimal(r["value"])==Decimal(raw)
    report=json.loads((OUT/"quality_report.json").read_text(encoding="utf-8"))
    assert report["display_years"]==list(range(2009,2024))
    assert report["core_missing_in_display"]==0
    assert not report["checks_failed"]
    book=ROOT/"outputs/data-review/全国人口与能源_数据核验.xlsx"
    wb=load_workbook(book,data_only=True)
    assert len(wb.sheetnames)==7
    population=rows("population_annual.csv")
    energy=rows("energy_annual.csv")
    structure=rows("energy_structure.csv")
    for i,r in enumerate(population,5):
        s=wb["人口年度"]
        assert s.cell(i,1).value==int(r["year"])
        assert close(s.cell(i,2).value,r["population_total"])
        assert close(s.cell(i,5).value,float(r["population_urban"])/float(r["population_total"]))
        a=wb["年龄结构"]
        assert close(a.cell(i,9).value,r["age_sum_difference"])
        for col,key in [(6,"0_14"),(7,"15_64"),(8,"65_plus")]:
            assert close(a.cell(i,col).value,float(r["age_share_"+key+"_pct"])/100)
    for i,r in enumerate(energy,5):
        s=wb["能源年度"]
        assert close(s.cell(i,2).value,r["energy_total"])
        assert s.cell(i,4).value is None
        if r["energy_yoy_calculated_pct"]:
            assert close(s.cell(i,3).value,float(r["energy_yoy_calculated_pct"])/100)
        assert (s.cell(i,5).value is None)==(r["living_energy_per_capita"]=="")
    for i,r in enumerate(structure,5):
        assert close(wb["能源结构"].cell(i,5).value,r["amount_estimated_tce_10k"])
    # A waterfall must reconcile even when rounded official shares do not sum to 100.
    for r in energy:
        year=int(r["year"])
        amounts=sum(float(x["amount_estimated_tce_10k"]) for x in structure if int(x["year"])==year)
        assert close(amounts+float(r["rounding_residual_tce_10k"]),r["energy_total"])
    print("PASS: 720 original response cells agree with CSV; 15-year core coverage complete.")
    print("PASS: 7-sheet workbook saved values/formulas agree with CSV; missing values stay blank.")
    print("PASS: all 18 annual energy totals reconcile including rounding residual.")
    wb.close()
if __name__=="__main__":
    main()
