import hashlib,unittest,json
from collections import Counter,defaultdict
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from matplotlib.backends.backend_agg import FigureCanvasAgg
from data_model import get_data,HR_SOURCE,SALES_SOURCE,ROOT,EDUCATION,DEPARTMENTS,age_band
from charts import render_component,render_dashboard,HR_SIZES,SALES_SIZES

class DashboardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d=get_data(); cls.hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in [HR_SOURCE,SALES_SOURCE]}
    @classmethod
    def tearDownClass(cls):
        for p,digest in cls.hashes.items(): assert hashlib.sha256(p.read_bytes()).hexdigest()==digest

    def test_hr_source_totals_and_corrected_education(self):
        h=self.d.hr(); self.assertEqual(h['total'],1470); self.assertEqual(self.d.unique_employee_ids,1375)
        self.assertEqual(h['age'],[97,229,325,297,522]); self.assertEqual(h['gender'],[882,588])
        self.assertEqual(h['education'],[170,282,572,398,48])
        self.assertEqual(h['marriage'],[470,673,327]); self.assertEqual(h['departments'],[437,799,162,9,46,17])
        self.assertEqual(h['changes'],[40,5,5])
        for key in ['age','gender','education','marriage','departments']:self.assertEqual(sum(h[key]),1470)
        self.assertEqual([age_band(x) for x in [18,24,25,29,30,34,35,39,40,60]],['18-24','18-24','25-29','25-29','30-34','30-34','35-39','35-39','40=<','40=<'])
        self.assertTrue(all(age_band(r[1])==r[7] for r in self.d.people))

    def test_september_against_source_cached_statistics(self):
        w=load_workbook(SALES_SOURCE,read_only=True,data_only=True)
        try:
            s=w['统计数据']; d=self.d.sales(9)
            for key,cell in [('count','I2'),('amount','K2'),('profit','M2'),('margin','N3'),('cost_total','G2'),('mom','L9')]:self.assertAlmostEqual(d[key],s[cell].value,places=8)
            for key,cells in [('regions',['E9','E10','E11','E12','E13','E14']),('couriers',['E18','E19','E20','E21','E22','E23']),('category',['K13','K14','K15']),('costs',['E27','E29','E30','E28'])]:
                np.testing.assert_allclose(d[key],[s[c].value for c in cells],rtol=1e-12)
            for i,p in enumerate(d['top3'],3):
                self.assertEqual(p['name'],s[f'X{i}'].value); self.assertAlmostEqual(p['amount'],s[f'Y{i}'].value); self.assertEqual(p['count'],s[f'Z{i}'].value)
            np.testing.assert_allclose(d['trend'],[s.cell(r,11).value for r in range(18,30)],rtol=1e-12)
        finally:w.close()

    def test_all_months_independent_aggregation(self):
        total_count=0; total_amount=0
        for month in range(1,13):
            d=self.d.sales(month); rows=[r for r in self.d.orders if r[9]==f'{month}月']
            self.assertEqual(d['count'],len(rows)); self.assertAlmostEqual(d['amount'],sum(r[4] for r in rows)/10000)
            self.assertAlmostEqual(d['profit'],sum(r[6] for r in rows)/10000)
            for key in ['regions','category','couriers']:self.assertEqual(sum(d[key]),len(rows))
            self.assertAlmostEqual(sum(d['costs']),d['cost_total']); self.assertAlmostEqual(d['margin'],d['profit']/d['amount'])
            sums=defaultdict(float); counts=Counter()
            for r in rows:sums[r[8]]+=r[4]; counts[r[8]]+=1
            expected=sorted(sums,key=lambda n:-sums[n])[:3]
            self.assertEqual([p['name'] for p in d['top3']],expected)
            for p in d['top3']:self.assertAlmostEqual(p['amount'],sums[p['name']]/10000); self.assertEqual(p['count'],counts[p['name']])
            if month==1:self.assertIsNone(d['mom']);self.assertIsNone(d['previous'])
            else:self.assertAlmostEqual(d['mom'],len(rows)/len(self.d.monthly[month-1])-1)
            total_count+=d['count'];total_amount+=d['amount']
        self.assertEqual(total_count,8564); self.assertAlmostEqual(total_amount,sum(r[4] for r in self.d.orders)/10000)
        for value in [0,13,None,'9',True]:
            with self.assertRaises(ValueError):self.d.sales(value)

    def test_render_all_138_components_and_13_dashboards(self):
        for kind,months,sizes in [('hr',[9],HR_SIZES),('sales',range(1,13),SALES_SIZES)]:
            for month in months:
                for i,size in enumerate(sizes):
                    fig=render_component(kind,i,month,self.d); canvas=FigureCanvasAgg(fig);canvas.draw()
                    self.assertEqual(canvas.get_width_height(),size)
                    renderer=canvas.get_renderer()
                    for text in fig.axes[0].texts:
                        box=text.get_window_extent(renderer)
                        self.assertGreaterEqual(box.x0,-2,(kind,month,i,text.get_text()))
                        self.assertLessEqual(box.x1,size[0]+2,(kind,month,i,text.get_text()))
                        self.assertGreaterEqual(box.y0,-2);self.assertLessEqual(box.y1,size[1]+2)
                    fig.clear()
                fig=render_dashboard(kind,month,self.d); canvas=FigureCanvasAgg(fig);canvas.draw()
                self.assertEqual(canvas.get_width_height(),(1200,650) if kind=='hr' else (1728,1092));fig.clear()

    def test_tk_controls_and_exports(self):
        from app import DashboardApp,create_root
        root=create_root();root.withdraw()
        try:
            app=DashboardApp(root);root.update()
            self.assertEqual(str(app.month_box['state']),'disabled')
            app.selector.current(1);app.selector.event_generate('<<ComboboxSelected>>');root.update()
            self.assertEqual(app.kind,'sales'); self.assertEqual(str(app.month_box['state']),'readonly')
            images=[]
            for month in [1,9,12]:
                app.month_box.current(month-1);app.month_box.event_generate('<<ComboboxSelected>>');root.update()
                self.assertEqual(app.month,month);images.append(hashlib.sha256(app.image.tobytes()).hexdigest())
                self.assertIn(f'年{month}月',app.status.get())
            self.assertEqual(len(set(images)),3)
            app.reset();self.assertEqual(app.month,9)
            path=app.save_current();self.assertTrue(path.is_relative_to(ROOT));self.assertTrue(path.exists())
            payload=json.loads(path.with_suffix('.json').read_text(encoding='utf8'));self.assertEqual(payload['数据']['count'],751)
            win=app.full_size();root.update();win.destroy()
            app.selector.current(0);app.select();self.assertEqual(app.kind,'hr')
        finally:root.destroy()

if __name__=='__main__':unittest.main(verbosity=2)
