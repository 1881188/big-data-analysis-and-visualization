"""Read-only workbook inventory. Never save input workbooks."""
from pathlib import Path
import json, warnings, zipfile, hashlib
from openpyxl import load_workbook
warnings.filterwarnings('ignore',category=UserWarning,module='openpyxl')
ROOT=Path(__file__).resolve().parent
result={}
for path in ROOT.parent.glob('第四章*.xlsx'):
    wb=load_workbook(path,data_only=False); cache=load_workbook(path,data_only=True)
    info={'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':{}}
    for ws in wb:
        raw=ws.title in ('202203人员基础信息','销售明细')
        data={'size':[ws.max_row,ws.max_column],'headers':[c.value for c in ws[1]],'cells':[],
              'print_area':str(ws.print_area),'merges':[str(r) for r in ws.merged_cells.ranges],
              'validations':str(ws.data_validations)}
        for row in ws:
            for c in row:
                if c.value is not None and (not raw or c.row==1 or (ws.title=='202203人员基础信息' and c.column>7)):
                    data['cells'].append({'cell':c.coordinate,'value':c.value,'cached':cache[ws.title][c.coordinate].value})
        info['sheets'][ws.title]=data
        print(path.name,ws.title,ws.max_row,ws.max_column,'headers',data['headers'])
        if raw: print('row2', [c.value for c in ws[2]])
    with zipfile.ZipFile(path) as z:
        info['xml']={n:z.read(n).decode('utf8') for n in z.namelist() if n.endswith('.xml') and n.startswith(('xl/charts/chart','xl/drawings/drawing','xl/slicer','xl/ctrlProps','xl/pivotTables/'))}
    result[path.name]=info; wb.close(); cache.close()
(ROOT/'源文件检查.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf8')
