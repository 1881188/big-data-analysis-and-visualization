"""Read source facts; never save or execute formulas/macros in original files."""
from pathlib import Path
from collections import Counter, defaultdict
from functools import lru_cache
import math
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
HR_SOURCE=ROOT.parent/'第四章 人力资源可视化看板.xlsx'
SALES_SOURCE=ROOT.parent/'第四章 销售看板参考.xlsx'
EDUCATION=['专科以下','专科','本科','硕士研究生','博士研究生']
DEPARTMENTS=['销售部','研发部','信息技术部','行政部','人力资源部','财务部']
AGE_LABELS=['18-24','25-29','30-34','35-39','40=<']
REGIONS=['华北','华南','东北','西北','西南','华东']
COURIERS=['顺丰','韵达','中通','申通','圆通','EMS']
CATEGORIES=['办公用品','家具产品','数码电子']
COSTS=['推广','人工','产品','其他']

def age_band(age):
    if age<18: raise ValueError('源定义未包含18岁以下人员')
    return AGE_LABELS[0 if age<25 else 1 if age<30 else 2 if age<35 else 3 if age<40 else 4]

class DashboardData:
    def __init__(self):
        wb=load_workbook(HR_SOURCE,read_only=True,data_only=True)
        try: self.people=list(wb.active.iter_rows(min_row=2,max_col=8,values_only=True))
        finally: wb.close()
        wb=load_workbook(SALES_SOURCE,read_only=True,data_only=True)
        try:
            self.orders=list(wb['销售明细'].iter_rows(min_row=2,max_col=10,values_only=True))
            self.cost_rows=list(wb['成本明细'].iter_rows(min_row=2,max_col=3,values_only=True))
            self.product_order=[r[0] for r in wb['统计数据'].iter_rows(min_row=2,max_row=347,min_col=19,max_col=19,values_only=True)]
        finally: wb.close()
        # Source COUNTA counts records. Duplicate identifiers do not authorize deduplication.
        self.unique_employee_ids=len({r[0] for r in self.people})
        if any(r[0] is None for r in self.people): raise ValueError('人员记录缺少编号')
        if len({r[1] for r in self.orders})!=len(self.orders): raise ValueError('订单单号重复，需重新确认销量口径')
        if any(r[9]!=f'{r[0].month}月' for r in self.orders): raise ValueError('订单日期与月份列不一致')
        years={r[0].year for r in self.orders}
        if len(years)!=1: raise ValueError('源数据含多个年份，请先明确年份筛选')
        self.year=next(iter(years))
        self.monthly={m:[r for r in self.orders if r[0].month==m] for m in range(1,13)}
        self.trend=[math.fsum(r[4] for r in self.monthly[m])/10000 for m in range(1,13)]

    def hr(self):
        counts=lambda column,labels: [sum(r[column]==v for r in self.people) for v in labels]
        ages=Counter(age_band(r[1]) for r in self.people)
        return dict(total=len(self.people),age=[ages[k] for k in AGE_LABELS],
                    gender=counts(2,['男','女']),marriage=counts(4,['单身','已婚','离异']),
                    education=counts(3,EDUCATION),departments=counts(5,DEPARTMENTS),
                    changes=counts(6,['入职','转入','转出']))

    def sales(self,month=9):
        if type(month) is not int or not 1<=month<=12: raise ValueError('月份必须为1至12的整数')
        rows=self.monthly[month]
        count=len(rows); prev=len(self.monthly[month-1]) if month>1 else None
        amount=math.fsum(r[4] for r in rows)/10000
        profit=math.fsum(r[6] for r in rows)/10000
        cost=[math.fsum(r[2] for r in self.cost_rows if r[0]==f'{month}月' and r[1]==c) for c in COSTS]
        products=defaultdict(list)
        for r in rows: products[r[8]].append(r)
        ranked=[dict(name=n,amount=math.fsum(r[4] for r in products[n])/10000,count=len(products[n])) for n in self.product_order]
        # Stable full-range descending order: fixes Excel's drifting RANK range and ties.
        ranked.sort(key=lambda p:-p['amount'])
        counter=lambda col,labels:[sum(r[col]==v for r in rows) for v in labels]
        return dict(month=month,year=self.year,count=count,previous=prev,mom=(count/prev-1) if prev else None,
                    amount=amount,profit=profit,margin=profit/amount if amount else None,
                    costs=cost,cost_total=math.fsum(cost),category=counter(7,CATEGORIES),
                    regions=counter(2,REGIONS),couriers=counter(3,COURIERS),top3=ranked[:3],trend=self.trend)

@lru_cache(maxsize=1)
def get_data(): return DashboardData()
