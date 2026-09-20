from pathlib import Path
import json, warnings, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
root = Path(__file__).resolve().parent
source = root.parent / "第三章 动态图表.xlsm"
book = load_workbook(source, data_only=False, read_only=False)
cached = load_workbook(source, data_only=True, read_only=True)
result = {}
for ws in book:
    entry = {"rows": ws.max_row, "columns": ws.max_column, "cells": [], "charts": []}
    if ws.title != "5 人力资源明细":
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    entry["cells"].append({"cell": cell.coordinate, "value": str(cell.value), "cached": cached[ws.title][cell.coordinate].value})
        for chart in ws._charts:
            entry["charts"].append(ET.tostring(chart.to_tree(),encoding="unicode"))
    else:
        entry["headers"] = [c.value for c in ws[1]]
        entry["sample"] = list(cached[ws.title].iter_rows(min_row=1,max_row=6,values_only=True))
        entry["distinct"] = {str(c): sorted({str(ws.cell(r,c).value) for r in range(2,ws.max_row+1)})[:60] for c in range(1,8)}
    result[ws.title] = entry
with zipfile.ZipFile(source) as z:
    result["_controls"] = {name: z.read(name).decode("utf-8",errors="replace") for name in z.namelist() if name.startswith(("xl/ctrlProps/","xl/slicerCaches/","xl/slicers/")) and name.endswith(".xml")}
(root / "源工作簿检查.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str),encoding="utf-8")
for name, entry in result.items():
    if name == "_controls": continue
    print("\n"+name)
    print(json.dumps({k:v for k,v in entry.items() if k != "charts"},ensure_ascii=False, default=str))
print("CONTROLS",json.dumps(result["_controls"],ensure_ascii=False))
