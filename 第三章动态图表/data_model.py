"""Read original cells, never save the workbook or execute VBA."""
from pathlib import Path
from collections import defaultdict
from functools import lru_cache
import warnings
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / '第三章 动态图表.xlsm'
NAMES = ['动态柱形图', '动态跑道图', '动态南丁格尔圆环图', '动态组合图',
         '透视表切片器', 'VBA动态玉玦图', '动态滑珠图']
EDUCATION = ['本科', '博士', '高中', '硕士', '专科']

class WorkbookData:
    def __init__(self, source=SOURCE):
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            book = load_workbook(source, data_only=True, read_only=True)
        try:
            def block(sheet, cells):
                return [[c.value for c in row] for row in book[sheet][cells]]
            self.regions = block('1 动态柱形图', 'C3:H3')[0]
            self.months = [r[0] for r in block('1 动态柱形图', 'B4:B9')]
            self.sales = block('1 动态柱形图', 'C4:H9')
            self.staff_names = [r[0] for r in block('2 动态跑道图', 'B4:B9')]
            self.staff = block('2 动态跑道图', 'C4:F9')
            self.traffic_names = block('3 动态南丁格尔圆环图', 'C3:G3')[0]
            self.traffic = block('3 动态南丁格尔圆环图', 'C4:G5')
            self.products = block('4 动态组合图', 'C3:F3')[0]
            self.amount, self.profit = block('4 动态组合图', 'C4:F5')
            self.margin = [p / a for p, a in zip(self.profit, self.amount)]
            self.people = list(book['5 人力资源明细'].iter_rows(min_row=2, values_only=True))
            self.departments = sorted({r[2] for r in self.people})
            self.jade_names = block('6 VBA动态玉玦图', 'C3:F3')[0]
            self.jade = block('6 VBA动态玉玦图', 'C4:F5')
            self.rates = block('7 动态滑珠图', 'C3:H8')
        finally:
            book.close()

    def salaries(self, departments):
        unknown = set(departments) - set(self.departments)
        if unknown:
            raise ValueError(f'未知部门: {unknown}')
        groups = defaultdict(list)
        for row in self.people:
            if row[2] in departments:
                groups[row[3]].append(row[6])
        return [sum(groups[k]) / len(groups[k]) if groups[k] else None for k in EDUCATION]

    def beads(self, mode, index):
        if mode not in (0, 1) or not 0 <= index < 6:
            raise ValueError('视角须为0或1，选项须为0至5')
        return (self.regions, [r[index] for r in self.rates]) if mode == 0 else (self.months, self.rates[index])

@lru_cache(maxsize=1)
def get_data():
    return WorkbookData()

def default_state(case):
    return [dict(index=4), dict(index=3), dict(index=0),
            dict(visible=(True, True, True)), dict(departments=('市场拓展',)),
            dict(index=1), dict(mode=0, index=0)][case].copy()
